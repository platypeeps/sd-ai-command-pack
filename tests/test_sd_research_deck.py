"""`sd-research-kit deck`: the status brief's deck.json laid out as PowerPoint (pack sd:5525).

The template `skills/sd-research-repo/templates/deck.json` is the fixture, so the
example the standard points at is the one the builder is proven against.
"""

from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
import unittest.mock
from pathlib import Path

from pptx import Presentation

from tests.clean_env import clean_environment

REPO_ROOT = Path(__file__).resolve().parents[1]
KIT = REPO_ROOT / "bin" / "sd-research-kit"
TEMPLATE = REPO_ROOT / "skills" / "sd-research-repo" / "templates" / "deck.json"


def load_deck_module():
    spec = importlib.util.spec_from_file_location("sd_research_deck", REPO_ROOT / "bin" / "sd_research_deck.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


deck_mod = load_deck_module()


def kit(*args, cwd):
    return subprocess.run([sys.executable, str(KIT), *args], cwd=cwd, capture_output=True, text=True, timeout=120,
                          env=clean_environment())


def texts(slide) -> str:
    return "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


def links(slide) -> list[str]:
    found = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                found += [run.hyperlink.address for run in para.runs if run.hyperlink.address]
    return found


class TemplateDeck(unittest.TestCase):
    def setUp(self):
        self.deck = json.loads(TEMPLATE.read_text(encoding="utf-8"))
        self.tmp = Path(tempfile.mkdtemp())

    def write(self, name, deck):
        path = self.tmp / name
        path.write_text(json.dumps(deck), encoding="utf-8")
        return path

    def test_the_template_validates(self):
        self.assertIs(deck_mod.check_deck(self.deck), self.deck)

    def test_deck_lays_out_this_repos_deck_json_into_the_dashboard_folder(self):
        (self.tmp / "30-brief").mkdir()
        self.write("30-brief/deck.json", self.deck)
        result = kit("deck", cwd=self.tmp)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("3 slides", result.stdout)
        prs = Presentation(str(self.tmp / "docs" / "dashboard" / "deck.pptx"))
        self.assertEqual(len(prs.slides), 3)
        title, first, second = prs.slides
        self.assertIn("STATUS  ·  OCTOBER 20, 2026", texts(title))
        self.assertIn("Example Revamp", texts(title))
        self.assertIn("A targeted revamp,", texts(first))
        self.assertIn("not a rebuild.", texts(first))
        self.assertIn("EXAMPLE REVAMP", texts(first))
        self.assertIn("we are here", texts(second))
        # The one call to action sits on the last slide only; the Details link on the slide that names it.
        self.assertEqual(links(title) + links(first), [])
        self.assertEqual(links(second), ["https://example.test/status", "https://example.test/"])

    def test_deck_without_a_deck_json_says_where_it_looks(self):
        result = kit("deck", cwd=self.tmp)
        self.assertEqual(result.returncode, 2)
        self.assertIn("30-brief/deck.json", result.stderr)

    def test_deck_refuses_an_invalid_deck_and_writes_nothing(self):
        (self.tmp / "30-brief").mkdir()
        self.deck["slides"][0]["cards"] *= 3
        self.write("30-brief/deck.json", self.deck)
        result = kit("deck", cwd=self.tmp)
        self.assertEqual(result.returncode, 1)
        self.assertIn("slides[0].cards: at most 6 cards", result.stderr)
        self.assertFalse((self.tmp / "docs").exists())

    def test_a_logo_resolves_beside_its_deck_json(self):
        from PIL import Image

        Image.new("RGB", (64, 16), "white").save(self.tmp / "logo.png")
        self.deck["theme"]["logo"] = "logo.png"
        loaded = deck_mod.load_deck(self.write("deck.json", self.deck))
        self.assertEqual(loaded["theme"]["logo"], str((self.tmp / "logo.png").resolve()))
        out = self.tmp / "out.pptx"
        deck_mod.build(loaded, out)
        pictures = [s for s in Presentation(str(out)).slides[1].shapes if s.shape_type == 13]
        self.assertEqual(len(pictures), 1)

    def test_a_failed_save_keeps_the_old_file_and_leaves_no_partial(self):
        out = self.tmp / "deck.pptx"
        out.write_bytes(b"old")
        with unittest.mock.patch("pptx.presentation.Presentation.save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                deck_mod.build(self.deck, out)
        self.assertEqual(out.read_bytes(), b"old")
        self.assertEqual(sorted(p.name for p in self.tmp.iterdir()), ["deck.pptx"])


class Refusals(unittest.TestCase):
    """Every refusal names the value's path; extreme and wrong-typed input refuses, never crashes."""

    def setUp(self):
        self.deck = json.loads(TEMPLATE.read_text(encoding="utf-8"))

    def refuses(self, said):
        with self.assertRaises(deck_mod.Invalid) as raised:
            deck_mod.check_deck(self.deck)
        self.assertIn(said, str(raised.exception))

    def test_unknown_and_missing_keys(self):
        self.deck["colour"] = "red"
        self.refuses("deck: unknown key colour")
        del self.deck["colour"], self.deck["date"]
        self.refuses("deck: missing date")

    def test_text_bounds_empty_and_em_dash(self):
        self.deck["slides"][0]["title"] = ""
        self.refuses("slides[0].title: a non-empty string")
        self.deck["slides"][0]["title"] = "x" * 10_000_000
        self.refuses("slides[0].title: 10000000 characters")
        self.deck["slides"][0]["title"] = "Works — mostly"
        self.refuses("em dash")

    def test_wrong_types_at_every_depth(self):
        for path, value, said in (
            (("slides",), {"a": 1}, "deck.slides: 1 to 3"),
            (("slides",), [], "deck.slides: 1 to 3"),
            (("slides", 0), [[["deep"]]], "slides[0]: an object"),
            (("slides", 1, "track"), [{}], "track: 2 to 5"),
            (("slides", 1, "track", 0, "state"), "later", "track[0].state: one of"),
            (("slides", 1, "cards", 0, "big"), 47, "cards[0].big: a non-empty string"),
            (("theme", "ink"), "#12345", "theme.ink: a #RRGGBB"),
            (("theme", "tones"), {"now": ["#FFFFFF"]}, "theme.tones.now"),
            (("theme", "tones"), {"loud": ["#FFFFFF", "#000000"]}, "theme.tones: unknown key loud"),
            (("theme", "cta", "url"), "javascript:alert(1)", "theme.cta.url"),
            (("date",), "2026-10-20T00:00", "deck.date"),
            (("date",), float("inf"), "deck.date"),
        ):
            with self.subTest(path=path):
                self.deck = json.loads(TEMPLATE.read_text(encoding="utf-8"))
                target = self.deck
                for key in path[:-1]:
                    target = target[key]
                target[path[-1]] = value
                self.refuses(said)

    def test_a_track_leaves_room_for_one_row_of_cards(self):
        self.deck["slides"][1]["cards"].append({"text": "a fourth"})
        self.refuses("slides[1].cards: at most 3 cards beside a track")

    def test_a_tone_needs_a_pill_and_a_foot_label_needs_a_foot(self):
        del self.deck["slides"][0]["cards"][0]["pill"]
        self.refuses("cards[0].tone: needs a pill")
        self.deck = json.loads(TEMPLATE.read_text(encoding="utf-8"))
        del self.deck["slides"][1]["foot"]
        self.refuses("foot_label: needs a foot")

    def test_unreadable_and_non_json_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            bad = Path(tmp) / "deck.json"
            bad.write_text("{", encoding="utf-8")
            with self.assertRaisesRegex(deck_mod.Invalid, "not JSON"):
                deck_mod.load_deck(bad)
            with self.assertRaisesRegex(deck_mod.Invalid, "cannot be read"):
                deck_mod.load_deck(Path(tmp) / "missing.json")


if __name__ == "__main__":
    unittest.main()
